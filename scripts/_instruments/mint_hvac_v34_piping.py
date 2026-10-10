"""SLICE 12e-1 -- MINT HVAC v34 = v33 + THE PIPING CATALOGUE (40 items, one data-only config, two live MS links).

Runs IN THE CONTAINER under the bench's python (it needs frappe for the ONE uid mint and, with --load, for
the loader / exporter). Two phases, both explicit:

  build   read the owner's workbook sheet "Piping" + the v33 asset, mint one item_uid per row through
          `csv_importer.mint_item_uid` (the ONE mint -- never inline), assemble the candidate payload
          (v33's items + 40 Piping items; v33's 9 configs + the Piping config; every other top-level key
          carried verbatim), validate it offline (config validator + the derived-rate consistency the
          loader will re-check), write it to --out. NO database write.
  load    take the pre-load HVAC snapshot, `loader.load_rate_master(path=<candidate>, replace=True)`,
          then `exporter.export_asset_text("HVAC")` -> the DB-exported v34 text (two exports of an
          unchanged DB are byte-identical) -> --canonical. Electrical is never touched.

The committed artefact is the DB export RE-ORDERED INTO THE FILE CONVENTION the HVAC series carries
(measured on v32 / v33, 2026-10-09): items in GLOBAL `item_uid` order (kinds interleaved), configs in
creation order with the new one LAST -- the exporter's own `kind asc, item_uid asc` / `category_id asc`
order is a different, content-equal ordering, and the cross-version pins compare ORDERED lists -- plus
the top-level `intentional_removals` key the exporter does not emit, carried from the previous asset.
Both are done by hand after `load` (the 12e-1 session record names the steps); the content of every
carried item and config is asserted byte-equal to v33 while doing so. Re-running `build` mints fresh
uids (the mint is random by design); the committed file is the record.

Owner rulings carried (2026-10-09): 40 items all per metre; sizes in the sheet's own spelling; copper
install flat 110 BCS; copper 31.7 install 220 (the sheet carries it); no rounding; MS 300 = 2 x MS 150
and MS 250 = MS 150 + MS 100 as LIVE links; one category at a time; every item carries its own
supply / install markup (ruling revised 2026-10-09 17:39); NO pricing (pipelines {}).

SLICE 12e-1b (owner 2026-10-10: "the rate master is missing the item name column. it should be kept
exactly as per the excel file. we cant drop it."; ruling "option a is ok"): EVERY PIPING ITEM CARRIES
`attributes.item_name` = ITS SHEET COLUMN A TEXT, VERBATIM -- the key is MANDATORY, never optional. The
Piping config declares it FIRST in `attribute_definitions` as a plain choice over the four sheet texts
(label "Item", NO `attributes_from_spec` -- the spec reader has no Piping rules, so spec mode would
store every item not_understood). It is the same shape Insulation's `type` definition carries (no
selector / panel flag), rendered where Insulation's column A renders: after brand and unit, before
Pipe type. The viewer and the rate file show an attribute ONLY when the config declares it, which is
why the definition and the key travel together.

  name    the third phase, HVAC v35 = v34 + that key: read the workbook's column A per row, stamp
          `item_name` on the base asset's 40 Piping items (matched by `source.row`, every other key
          untouched, uids KEPT -- nothing is minted), insert the definition first in the Piping config,
          validate offline exactly as `build` does, write the candidate. NO database write; `load` then
          runs as before, and the same two hand steps follow the export.
`build` (a fresh mint) carries the key and the definition too, so a future re-mint cannot drop it.
"""
import argparse, copy, json, os, sys

PIPE_TYPE_BY_ITEM = {
    "Refrigerant Piping/Copper Piping": "Copper",
    "Chilled Water Piping / MS Piping": "MS",
    "PVC Pipe": "PVC",
    "CPVC Pipe": "CPVC",
}
KIND = "hvac_piping_item"
CATEGORY_ID = "hvac_piping"
SHEET = "Piping"
# header -> column index (0-based), verified against the sheet's row 1 on 2026-10-09
COLS = {"item": 0, "unit": 1, "size_mm": 2, "size_in": 3, "supply_markup": 4, "install_markup": 5,
        "bcs_pipe": 6, "bcs_accessories": 7, "cost_supply": 8, "cost_install": 9, "boq_supply": 10, "boq_install": 11}
HEADERS = ["Item", "Unit", "Pipe Size/ Dia (mm)", "Pipe Size/ Dia (inches)", "Supply Mark up", "Install Mark up",
           "BCS Pipe", "BCS Accessories", "Total BCS Supply", "BCS Installation", "BoQ Supply", "BoQ Installation"]


def read_sheet(path):
    import openpyxl
    ws = openpyxl.load_workbook(path, data_only=True)[SHEET]
    rows = [r for r in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True) if any(v is not None for v in r)]
    if [str(h).strip() for h in rows[0][:12]] != HEADERS:
        raise SystemExit("sheet header differs from the verified shape: %r" % (rows[0][:12],))
    out = []
    for excel_row, r in enumerate(rows[1:], start=2):
        raw = r[COLS["item"]]
        if raw is None or not str(raw).strip():
            raise SystemExit("row %d: column A (Item) is blank -- every Piping item needs its sheet name (12e-1b)" % excel_row)
        # 12e-1b: the name is kept VERBATIM (whitespace and slashes as written); only a trailing space goes
        item_name = str(raw).rstrip(" ")
        item = item_name.strip()
        if item not in PIPE_TYPE_BY_ITEM:
            raise SystemExit("row %d: unknown Item text %r" % (excel_row, item))
        num = lambda k: r[COLS[k]]
        for k in ("size_mm", "supply_markup", "install_markup", "cost_supply", "cost_install", "bcs_pipe", "bcs_accessories", "boq_supply", "boq_install"):
            if not isinstance(num(k), (int, float)):
                raise SystemExit("row %d: %s is not a number: %r" % (excel_row, k, num(k)))
        out.append({
            "excel_row": excel_row, "item_name": item_name, "pipe_type": PIPE_TYPE_BY_ITEM[item], "unit": str(r[COLS["unit"]]).strip(),
            "size_mm": float(num("size_mm")), "supply_markup": float(num("supply_markup")), "install_markup": float(num("install_markup")),
            "cost_supply": float(num("cost_supply")), "cost_install": float(num("cost_install")),
            # 12e-2: the sheet's BCS Pipe (G), BCS Accessories (H) and the two BoQ figures (K, L) -- read, verified, never computed
            "bcs_pipe": float(num("bcs_pipe")), "bcs_accessories": float(num("bcs_accessories")),
            "boq_supply": float(num("boq_supply")), "boq_install": float(num("boq_install")),
        })
    if len(out) != 40:
        raise SystemExit("expected 40 data rows, read %d" % len(out))
    return out


def item_name_definition(rows):
    """12e-1b: the FIRST attribute definition of the Piping config -- a plain choice over the sheet's column A
    texts, verbatim, in sheet order (first seen). The shape Insulation's `type` definition carries."""
    values = []
    for r in rows:
        if r["item_name"] not in values:
            values.append(r["item_name"])
    return {"id": "item_name", "label": "Item", "type": "choice", "values": values,
            "note": "The sheet's own Item text (column A), verbatim -- every Piping item carries it (owner 2026-10-10: "
                    "'it should be kept exactly as per the excel file. we cant drop it.'). A plain attribute, not a spec-read one."}


def _validate_candidate(config, items):
    """The offline checks both minting phases run: the loader's own validators + the derived-rate consistency."""
    from nirmaan_stack.services.boq_rate_master import config_validation
    from nirmaan_stack.services.boq_rate_master.loader import _validate_items, _validate_one_config
    _validate_one_config(config, "category_configs[hvac_piping]")
    config_validation._validate_config(config)
    _validate_items(items)
    by_uid = {it["item_uid"]: {"rates": it["rates"]} for it in items}
    bad = config_validation.derived_rate_updates([config], by_uid)
    if bad:
        raise SystemExit("derived cells disagree with their base in the candidate: %r" % bad)
    missing = [it["item_uid"] for it in items if it["kind"] == KIND and not str(it["attributes"].get("item_name") or "").strip()]
    if missing:
        raise SystemExit("12e-1b: every Piping item carries item_name; missing on %r" % missing)


def name_items(workbook, base, out):
    """12e-1b `name`: v35 = base (v34) + `attributes.item_name` on the 40 Piping items + the definition, first.
    Matched by `source.row`; uids kept; nothing else touched; NO database write."""
    import frappe
    os.chdir("/workspace/development/frappe-bench/sites")
    frappe.init(site="localhost"); frappe.connect()
    rows = read_sheet(workbook)
    by_row = {r["excel_row"]: r for r in rows}
    payload = json.load(open(base, encoding="utf-8"))
    items, stamped = [], 0
    for it in payload["items"]:
        if it["kind"] != KIND:
            continue
        r = by_row.get(it["source"]["row"])
        if r is None or it["source"]["sheet"] != SHEET:
            raise SystemExit("Piping item %s has no sheet row to read its name from: %r" % (it["item_uid"], it["source"]))
        if r["pipe_type"] != it["attributes"]["pipe_type"] or float(r["size_mm"]) != float(it["attributes"]["size_mm"]):
            raise SystemExit("Piping item %s does not match sheet row %d" % (it["item_uid"], it["source"]["row"]))
        if "item_name" in it["attributes"]:
            raise SystemExit("Piping item %s already carries item_name" % it["item_uid"])
        it["attributes"] = dict([("item_name", r["item_name"])] + list(it["attributes"].items()))   # the name FIRST, as declared
        items.append(it); stamped += 1
    if stamped != 40:
        raise SystemExit("expected to stamp 40 Piping items, stamped %d" % stamped)
    config = next(c for c in payload["category_configs"] if c["category_id"] == CATEGORY_ID)
    if any(d.get("id") == "item_name" for d in config["attribute_definitions"]):
        raise SystemExit("the Piping config already declares item_name")
    config["attribute_definitions"] = [item_name_definition(rows)] + config["attribute_definitions"]
    _validate_candidate(config, items)
    json.dump(payload, open(out, "w", encoding="utf-8"), indent=1, ensure_ascii=False)
    print("candidate written:", out, "| items", len(payload["items"]), "| configs", len(payload["category_configs"]),
          "| named", stamped, "| definitions", [d["id"] for d in config["attribute_definitions"]])
    frappe.destroy()


def build(workbook, base, out):
    import frappe
    os.chdir("/workspace/development/frappe-bench/sites")
    frappe.init(site="localhost"); frappe.connect()
    from nirmaan_stack.services.boq_rate_master import csv_importer, config_validation
    from nirmaan_stack.services.boq_rate_master.loader import _validate_items, _validate_one_config
    rows = read_sheet(workbook)
    v33 = json.load(open(base, encoding="utf-8"))
    taken = {it["item_uid"] for it in v33["items"]}
    taken |= {(r["item_uid"] or "").strip() for r in frappe.get_all("BoQ Rate Master Item", filters={"discipline": "HVAC"}, fields=["item_uid"])}
    items = []
    for r in rows:
        uid = csv_importer.mint_item_uid("HVAC", taken)     # the ONE mint
        items.append({
            "kind": KIND, "brand": None, "unit": r["unit"],
            "attributes": {"item_name": r["item_name"], "pipe_type": r["pipe_type"], "size_mm": r["size_mm"]},   # 12e-1b: the name, first
            "rates": {"cost_supply": r["cost_supply"], "cost_install": r["cost_install"],
                      "supply_markup": r["supply_markup"], "install_markup": r["install_markup"]},
            "source": {"sheet": SHEET, "row": r["excel_row"]},
            "item_uid": uid,
        })
    by_ms = {it["attributes"]["size_mm"]: it["item_uid"] for it in items if it["attributes"]["pipe_type"] == "MS"}
    ms150, ms100, ms300, ms250 = by_ms[150.0], by_ms[100.0], by_ms[300.0], by_ms[250.0]
    link = lambda uid, mult: {"from": {"item_uid": uid, "rate_key": "__KEY__"}, "multiplier": mult, "constant": 0.0}
    def links(terms):
        return {k: [dict(t, **{"from": {"item_uid": t["from"]["item_uid"], "rate_key": k}}) for t in terms] for k in ("cost_supply", "cost_install")}
    derived = {ms300: links([link(ms150, 2.0)]), ms250: links([link(ms150, 1.0), link(ms100, 1.0)])}
    sizes = sorted({it["attributes"]["size_mm"] for it in items})
    config = {
        "discipline": "HVAC",
        "category_id": CATEGORY_ID,
        "category_display": "Piping",
        "item_kinds": [KIND],
        "attribute_definitions": [
            item_name_definition(rows),   # 12e-1b: the sheet's Item text, FIRST
            {"id": "pipe_type", "label": "Pipe type", "type": "choice", "values": ["Copper", "MS", "PVC", "CPVC"]},
            {"id": "size_mm", "label": "Pipe size / dia (mm)", "type": "number_choice", "values": sizes},
        ],
        "pipelines": {},
        "derived_rates": derived,
        "notes": "SLICE 12e-1 (owner, 2026-10-09): the Piping CATALOGUE only -- 40 per-metre items from the HVAC workbook's "
                 "Piping sheet (copper 13, MS 13, PVC 8, CPVC 6), each with the sheet's own BCS supply / BCS install / supply "
                 "markup / install markup, sizes in the sheet's own spelling, no inch pre-computation. DATA-ONLY: no pipelines, so "
                 "it can never price a BoQ row and the calculator shows 'coming soon' (12e-2 brings the rules). The two MS links are "
                 "LIVE: MS 300 = 2 x MS 150 and MS 250 = MS 150 + MS 100 on both cost columns, recomputed on every write of the base.",
    }
    _validate_candidate(config, items)   # the loader's own validators + the post-load consistency check, offline
    payload = copy.deepcopy(v33)
    payload["items"] = v33["items"] + items
    payload["category_configs"] = v33["category_configs"] + [config]
    json.dump(payload, open(out, "w", encoding="utf-8"), indent=1, ensure_ascii=False)
    print("candidate written:", out, "| items", len(payload["items"]), "| configs", len(payload["category_configs"]),
          "| MS150", ms150, "MS100", ms100, "MS300", ms300, "MS250", ms250)
    frappe.destroy()


# ---------------------------------------------------------------------------------------------------------
# SLICE 12e-2 (owner option "1", 2026-10-10 10:58; rulings "option b, yes" / Q14 / Q14a / Q17 / Q11 / Q16 / Q15
# / Q10 / Q19 / Q18) -- PHASE `price`: HVAC v36 = v35 + PIPING PRICES IN THE CALCULATOR AND ON A BoQ ROW.
#
#   (a) cost_supply on the 40 Piping items = the sheet's BCS Pipe (column G), VERIFIED FIRST: for all 40 rows
#       ROUNDUP(G x (1 + the family's accessories)) must equal the sheet's Total BCS Supply (column I, what v35
#       stores), or the phase stops naming the rows. cost_install, both markups, item_name, pipe_type, size_mm
#       and the uids are untouched. The two MS links (derived_rates) now hold on column G (verified).
#   (b) FOUR new `hvac_pricing_input` items in the shape of the seven HVAC inputs ("Piping accessories - Copper"
#       0.30; "- MS", "- PVC", "- CPVC" 0.60), uids through the ONE mint.
#   (c) the Piping config becomes an ITEM-LIST config in Insulation's shape: families keyed by pipe_type, one
#       per-metre unit block each, the P5 pipeline (rate_ref to the family's input -> match -> BCS pipe ->
#       x (1 + accessories) -> ROUNDUP -> x (1 + supply markup) -> ROUNDUP; install: x (1 + install markup)
#       -> ROUNDUP), the P1 / P2 / P3 / P4 rules as config keys. `item_name` keeps its definition for the Rate
#       Master (selector: false, panel: false -- ADP's flags) and is never asked of the model nor shown on the
#       panel: the model's questions are `list_spec.attribute_definitions` alone.
# Offline validation exactly as `build` / `name` (the loader's validators + the derived-rate consistency +
# the asset check for item_name); NO database write. `load` then runs as before; the same two hand steps
# (series order, intentional_removals) follow the export.
# ---------------------------------------------------------------------------------------------------------
ACCESSORIES = {"Copper": 0.3, "MS": 0.6, "PVC": 0.6, "CPVC": 0.6}      # the sheet's column H, per family
INPUT_IDS = {"Copper": "piping_accessories_copper", "MS": "piping_accessories_ms",
             "PVC": "piping_accessories_pvc", "CPVC": "piping_accessories_cpvc"}
INPUT_KIND = "hvac_pricing_input"
# the length spellings Insulation's unit table carries (12c), "mts" first so the calculator's picker reads "mts"
LENGTH_SPELLINGS = ["mts", "rmt", "rm", "mtr", "m", "metre", "meter", "metres", "meters", "mt", "rmtr", "r.mt", "rmt.",
                    "running metre", "running meter", "r.m", "lm", "mtrs", "rmts"]


def _roundup(x):
    import math
    return math.ceil(round(x, 9))


def piping_input_items(taken):
    """The four accessories inputs, in the shape of the seven HVAC inputs (unit `factor`, one `rates.factor`)."""
    from nirmaan_stack.services.boq_rate_master import csv_importer
    out = []
    for fam in ("Copper", "MS", "PVC", "CPVC"):
        uid = csv_importer.mint_item_uid("HVAC", taken)     # the ONE mint
        taken.add(uid)
        out.append({
            "kind": INPUT_KIND, "brand": None, "unit": "factor",
            "attributes": {
                "item": INPUT_IDS[fam],
                "name": "Piping accessories - %s" % fam,
                "remarks": "The accessories share on a %s pipe's BCS cost (the sheet's 'BCS Accessories' column): "
                           "supply = ROUNDUP(BCS pipe x (1 + this)) before the supply markup. Read by the %s per-metre "
                           "pipeline only." % (fam, fam),
                "shared_by": "piping",
            },
            "rates": {"factor": ACCESSORIES[fam]},
            "source": {"sheet": "Pricing Inputs", "row": 0},
            "item_uid": uid,
        })
    return out


def _supply_pipeline(fam):
    label = "Piping accessories - %s" % fam
    return {"output": ["supply"], "steps": [
        {"step": "rate_ref", "ref": {"kind": INPUT_KIND, "item": INPUT_IDS[fam]}, "target": "factor", "result": "acc",
         "explain": "pricing input: %s (factor)" % INPUT_IDS[fam]},
        {"step": "match_master_row", "params": {"kind": KIND}, "explain": "match the pipe SKU (pipe type and size)"},
        {"step": "component", "name": "pipe", "target": "cost_supply", "formula": "base",
         "explain": "BCS pipe: the pipe's own cost from the sheet; the accessories share (the Pricing Input '%s') and the "
                    "supply markup are added below" % label},
        {"step": "sum_components", "result": "bcs_pipe", "explain": "BCS pipe"},
        {"step": "scale", "target": "bcs_pipe", "result": "bcs_supply", "params": {"acc_from_ctx": "acc"}, "formula": "base*(1+acc)",
         "explain": "BCS pipe x (1 + the accessories share)"},
        {"step": "roundup", "target": "bcs_supply", "params": {"digits": 0}, "explain": "ROUNDUP(BCS supply, 0)"},
        {"step": "scale", "target": "bcs_supply", "result": "supply", "params": {"m_from_ctx": "supply_markup"}, "formula": "base*(1+m)",
         "explain": "BCS supply x (1 + the SKU's supply markup)"},
        {"step": "roundup", "target": "supply", "params": {"digits": 0}, "explain": "ROUNDUP(supply, 0)"},
    ]}


def _install_pipeline():
    return {"output": ["install"], "steps": [
        {"step": "match_master_row", "params": {"kind": KIND}, "explain": "match the pipe SKU (pipe type and size)"},
        {"step": "component", "name": "install", "target": "cost_install", "formula": "base",
         "explain": "BCS install: the pipe's own installation cost from the sheet"},
        {"step": "sum_components", "result": "bcs_install", "explain": "BCS install"},
        {"step": "scale", "target": "bcs_install", "result": "install", "params": {"m_from_ctx": "install_markup"}, "formula": "base*(1+m)",
         "explain": "BCS install x (1 + the SKU's install markup)"},
        {"step": "roundup", "target": "install", "params": {"digits": 0}, "explain": "ROUNDUP(install, 0)"},
    ]}


def piping_pricing_config(base_config):
    """The v36 Piping config: v35's config + matching_mode item_list + the list_spec (the model's three questions,
    the pricing block with the P1-P5 rules) + ADP's two flags on the `item_name` definition. Everything else
    (item_kinds, derived_rates, category_display, the three definitions' content) is carried verbatim."""
    cfg = copy.deepcopy(base_config)
    defs = cfg["attribute_definitions"]
    assert [d["id"] for d in defs] == ["item_name", "pipe_type", "size_mm"], defs
    defs[0] = dict(defs[0], selector=False, panel=False)   # never asked of the model, never on the panel (ADP's shape)
    families = {}
    for fam in ("Copper", "MS", "PVC", "CPVC"):
        families[fam] = {"needs": [], "units": {"length": {"needs": ["size_mm"],
                                                           "pipelines": {"supply": _supply_pipeline(fam), "install": _install_pipeline()}}}}
    cfg["matching_mode"] = "item_list"
    cfg["list_spec"] = {
        "family_attribute_id": "pipe_type",
        "attribute_definitions": [
            {"id": "pipe_type", "label": "Pipe type", "type": "choice", "values": ["Copper", "MS", "PVC", "CPVC"],
             "note": "Which pipe the row buys: refrigerant / copper piping is Copper; chilled-water MS piping is MS; PVC and "
                     "CPVC as written. (The model instruction for GI, uPVC, HDPE, SS and 'cannot tell' is written in 12e-4.)"},
            {"id": "size_mm", "label": "Pipe size", "type": "text",
             "note": "The pipe's nominal size AS THE ROW WRITES IT, in mm or in inches (5/8\", 1-1/4\", 2 inch). Never converted."},
            {"id": "pipe_class", "label": "Class / wall thickness", "type": "text",
             "note": "A class, wall thickness, schedule or 'seamless' the row states for the pipe, copied as written. Recorded "
                     "and shown; it never changes the price (the sheet has one rate per size)."},
        ],
        "second_opinion": False,
        "pricing": {
            "kind": KIND,
            "label_attr": "pipe_type",
            "unit_class_attr": "unit_class",
            "unit_classes": {"length": LENGTH_SPELLINGS},
            "unit_words": {"length": "metre"},
            "unit_refusal": "unit '{unit}' is not a length unit",
            "numbers": {"size_mm": {"from": ["size_mm"], "name": "pipe size", "unit": "mm", "inches": True, "inch_mm_alt": 25,
                                    "typed_entry": "one_size"}},
            "ladders": ["size_mm"],
            "match_attrs": ["size_mm"],
            "choice_attrs": [],
            "size_match": {"dp": [2, 1], "near": 0.1, "below_smallest": "smallest"},
            "family_text": {
                "map": [
                    {"from": "GI", "to": "MS", "rule": "Q16 GI is priced as MS (owner 2026-10-10)"},
                    {"from": "uPVC", "to": "PVC", "rule": "Q16 uPVC is priced as PVC (owner 2026-10-10)"},
                    {"from": "HDPE", "to": "PVC", "rule": "Q16 HDPE is priced as PVC (owner 2026-10-10)"},
                ],
                "refuse": [
                    {"from": "SS", "refuse": "No SKU in the catalogue for SS pipe - price this row by hand",
                     "rule": "Q16 SS refuses (owner 2026-10-10)"},
                ],
                "line": "BoQ says {from} -> priced as {to} (owner rule)",
                "from_row": {"rule": "Q15 names that all land on one type price as that type; two different types with none chosen refuse (owner 2026-10-09)",
                             "refuse": "two pipe types are named in this row ({types}) - pick the type"},
            },
            "families": families,
            "panel_controls": {"pipe_type": "dropdown_or_other", "size_mm": "dropdown_or_other"},
            "panel_notes": {
                "pipe_type": "Pipe type as the BoQ writes it",
                "size_mm": "Type the pipe size the BoQ states, in mm or in inches (5/8\", 1-1/4\", 2 inch). For an NB size type the number.",
            },
            "panel_optional": ["pipe_class"],
            "read_notes": [
                {"families": ["Copper", "MS", "PVC", "CPVC"], "from_attr": "pipe_class", "pattern": "^\\s*(\\S.*?)\\s*$",
                 "line": "{match} stated -> priced at the one {family} rate (the sheet has no class rates)"},
            ],
        },
    }
    cfg["notes"] = (base_config.get("notes") or "") + (
        " SLICE 12e-2 (owner option 1, 2026-10-10): the PRICING RULES, as an item-list config in Insulation's shape -- the "
        "category is live on the calculator AND on a BoQ row the day its rules run (owner S1: no staging switch). "
        "cost_supply is the sheet's BCS Pipe (column G); supply = ROUNDUP(ROUNDUP(BCS pipe x (1 + the family's accessories "
        "Pricing Input)) x (1 + the SKU's supply markup)); install = ROUNDUP(BCS install x (1 + the SKU's install markup)). "
        "The BoQ figures are the sheet's columns K / L unchanged. item_name is never asked of the model and never on the panel.")
    return cfg


def price_items(workbook, base, out):
    """12e-2 `price`: v36 = base (v35) + cost_supply = BCS Pipe on the 40 Piping items (verified against column I
    first) + the four accessories inputs + the item-list Piping config. Matched by `source.row`; uids kept; NO
    database write."""
    import frappe
    os.chdir("/workspace/development/frappe-bench/sites")
    frappe.init(site="localhost"); frappe.connect()
    rows = read_sheet(workbook)
    by_row = {r["excel_row"]: r for r in rows}
    payload = json.load(open(base, encoding="utf-8"))
    # (a) VERIFY FIRST: ROUNDUP(G x (1 + accessories)) == I on every row, and the MS links hold on G
    bad = []
    for r in rows:
        acc = ACCESSORIES[r["pipe_type"]]
        if _roundup(r["bcs_pipe"] * (1 + acc)) != r["cost_supply"] or abs(r["bcs_accessories"] - acc) > 1e-9:
            bad.append((r["excel_row"], r["pipe_type"], r["size_mm"], r["bcs_pipe"], acc, r["cost_supply"]))
    if bad:
        raise SystemExit("AC2(a): ROUNDUP(G x (1 + accessories)) != I on %d row(s): %r" % (len(bad), bad))
    g = {(r["pipe_type"], r["size_mm"]): r["bcs_pipe"] for r in rows}
    if g[("MS", 300.0)] != 2 * g[("MS", 150.0)] or g[("MS", 250.0)] != g[("MS", 150.0)] + g[("MS", 100.0)]:
        raise SystemExit("AC2: the MS links do not hold on column G: %r" % {k: v for k, v in g.items() if k[0] == "MS"})
    items, stamped = [], 0
    for it in payload["items"]:
        if it["kind"] != KIND:
            continue
        r = by_row.get(it["source"]["row"])
        if r is None or it["source"]["sheet"] != SHEET:
            raise SystemExit("Piping item %s has no sheet row: %r" % (it["item_uid"], it["source"]))
        if r["pipe_type"] != it["attributes"]["pipe_type"] or float(r["size_mm"]) != float(it["attributes"]["size_mm"]):
            raise SystemExit("Piping item %s does not match sheet row %d" % (it["item_uid"], it["source"]["row"]))
        if it["rates"]["cost_supply"] != r["cost_supply"]:
            raise SystemExit("Piping item %s: v35 cost_supply %r is not the sheet's column I %r" % (it["item_uid"], it["rates"]["cost_supply"], r["cost_supply"]))
        it["rates"] = dict(it["rates"], cost_supply=float(r["bcs_pipe"]))   # the one key that moves; order kept
        items.append(it); stamped += 1
    if stamped != 40:
        raise SystemExit("expected to re-base 40 Piping items, did %d" % stamped)
    # (b) the four inputs
    taken = {it["item_uid"] for it in payload["items"]}
    taken |= {(x["item_uid"] or "").strip() for x in frappe.get_all("BoQ Rate Master Item", filters={"discipline": "HVAC"}, fields=["item_uid"])}
    inputs = piping_input_items(taken)
    payload["items"] = payload["items"] + inputs
    # (c) the config
    idx = next(i for i, c in enumerate(payload["category_configs"]) if c["category_id"] == CATEGORY_ID)
    config = piping_pricing_config(payload["category_configs"][idx])
    payload["category_configs"][idx] = config
    _validate_candidate(config, items + inputs)
    from nirmaan_stack.services.boq_rate_master import extraction
    assert extraction.has_runnable_pricing_rules(config) and extraction.config_is_eligible(config), "the rules must run"
    assert [d["id"] for d in extraction.build_items_spec(config)["attribute_definitions"]] == ["pipe_type", "size_mm", "pipe_class"]
    assert "item_name" not in [d["id"] for d in extraction.build_attribute_defs(config)]
    json.dump(payload, open(out, "w", encoding="utf-8"), indent=1, ensure_ascii=False)
    print("candidate written:", out, "| items", len(payload["items"]), "| configs", len(payload["category_configs"]),
          "| re-based", stamped, "| inputs", [i["item_uid"] for i in inputs])
    frappe.destroy()


def load(candidate, canonical):
    import frappe
    os.chdir("/workspace/development/frappe-bench/sites")
    frappe.init(site="localhost"); frappe.connect()
    frappe.set_user("Administrator")
    from nirmaan_stack.services.boq_rate_master import exporter, loader
    # AC8: the pre-load HVAC snapshot
    pre_payload, pre_text = exporter.export_asset_text("HVAC")
    snap = exporter.write_snapshot("HVAC", pre_text, pre_payload)
    frappe.db.commit()
    print("pre-load snapshot:", snap, "| items", len(pre_payload["items"]), "| configs", len(pre_payload["category_configs"]))
    summary = loader.load_rate_master(path=candidate, replace=True)
    frappe.db.commit()
    print("LOAD SUMMARY:", json.dumps(summary, default=str))
    payload, text = exporter.export_asset_text("HVAC")
    open(canonical, "w", encoding="utf-8").write(text)
    print("canonical export written:", canonical, "| items", len(payload["items"]), "| configs", len(payload["category_configs"]))
    frappe.destroy()


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("phase", choices=["build", "name", "price", "load"])
    ap.add_argument("--workbook"); ap.add_argument("--base"); ap.add_argument("--out")
    ap.add_argument("--candidate"); ap.add_argument("--canonical")
    a = ap.parse_args()
    if a.phase == "build":
        build(a.workbook, a.base, a.out)
    elif a.phase == "name":
        name_items(a.workbook, a.base, a.out)
    elif a.phase == "price":
        price_items(a.workbook, a.base, a.out)   # 12e-2
    else:
        load(a.candidate, a.canonical)
